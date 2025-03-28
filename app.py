from flask import Flask, request, Response
from twilio.twiml.messaging_response import MessagingResponse
from openai import OpenAI
import time
import os
from twilio.rest import Client as TwilioClient
import requests
from io import BytesIO
import fitz  # PyMuPDF para PDFs
from docx import Document  # Para Word
import openpyxl  # Para Excel
import gspread
from google.oauth2.service_account import Credentials
from datetime import datetime

app = Flask(__name__)

# Configuração da API OpenAI
client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY"),
    default_headers={"OpenAI-Beta": "assistants=v2"}
)

# Configuração da API Twilio
twilio_client = TwilioClient(
    os.getenv("TWILIO_ACCOUNT_SID"),
    os.getenv("TWILIO_AUTH_TOKEN")
)

ASSISTANT_ID = "asst_mlwRF5Byw4b4gqlYz9jvJtwV"
ultima_interacao = {}
user_threads = {}

# Conexão com Google Sheets
scope = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive"
]
creds = Credentials.from_service_account_file('credentials.json', scopes=scope)
gclient = gspread.authorize(creds)
sheet = gclient.open("Contatos WhatsApp Camila I.A").sheet1

# Funções auxiliares
def extrair_texto_pdf(conteudo):
    pdf_documento = fitz.open(stream=conteudo, filetype="pdf")
    texto = ""
    for pagina in pdf_documento:
        texto += pagina.get_text()
    return texto.strip()

def extrair_texto_word(conteudo):
    documento = Document(BytesIO(conteudo))
    return "\n".join([paragrafo.text for paragrafo in documento.paragraphs])

def extrair_texto_excel(conteudo):
    workbook = openpyxl.load_workbook(BytesIO(conteudo), data_only=True)
    sheet = workbook.active
    texto = ""
    for row in sheet.iter_rows(values_only=True):
        linha = [str(cell) for cell in row if cell is not None]
        texto += ' '.join(linha) + '\n'
    return texto.strip()

@app.route('/bot', methods=['POST'])
def whatsapp_reply():
    print("📩 Requisição recebida no /bot — retorno rápido")

    # ✅ Passo 1: resposta imediata para evitar erro 502/11200
    resposta_xml = "<?xml version='1.0' encoding='UTF-8'?><Response></Response>"
    return Response(resposta_xml, mimetype='text/xml')

    # ⛔️ O restante do processamento com OpenAI, mídia, planilhas etc
    # será feito no próximo passo, fora dessa função ou em background
    # (Aqui, temporariamente, ele está sendo interrompido)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)
